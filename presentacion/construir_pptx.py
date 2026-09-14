"""
presentacion/construir_pptx.py
Genera la presentación del Proyecto 2 (Miembro D).

Uso:
    python presentacion/construir_pptx.py

Produce `presentacion/proyecto2-eda-cervical.pptx`: 14 diapositivas en 16:9 que
siguen la misma paleta que el informe, reutilizando las figuras ya exportadas a
`figures/`. Si falta alguna figura, la diapositiva se genera con un marcador
visible en lugar de fallar, para que el problema no pase inadvertido.

Todas las cifras provienen de los notebooks 01 a 05 y coinciden con el informe.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
FIGURAS = RAIZ / "figures"
SALIDA = Path(__file__).resolve().parent / "proyecto2-eda-cervical.pptx"

# Paleta compartida con el informe y con src/plots.py
OSCURO = RGBColor(0x26, 0x46, 0x53)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
NARANJA = RGBColor(0xE7, 0x6F, 0x51)
GRIS = RGBColor(0x5B, 0x6B, 0x73)
GRIS_CLARO = RGBColor(0x9C, 0xA3, 0xAF)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_SUAVE = RGBColor(0xF4, 0xF6, 0xF7)

ANCHO = Inches(13.333)
ALTO = Inches(7.5)
FUENTE = "Segoe UI"

_contador = {"n": 0}


# ---------------------------------------------------------------------------
# Primitivas de maquetado
# ---------------------------------------------------------------------------
def _texto(marco, lineas, size=16, color=GRIS, bold_first=False, space=10):
    """Rellena un cuadro de texto con viñetas. `lineas` admite (texto, negrita)."""
    marco.word_wrap = True
    primero = True
    for linea in lineas:
        contenido, negrita = linea if isinstance(linea, tuple) else (linea, False)
        parrafo = marco.paragraphs[0] if primero else marco.add_paragraph()
        primero = False
        parrafo.space_after = Pt(space)
        parrafo.line_spacing = 1.18
        corrida = parrafo.add_run()
        corrida.text = contenido
        corrida.font.size = Pt(size)
        corrida.font.name = FUENTE
        corrida.font.color.rgb = OSCURO if (negrita or bold_first) else color
        corrida.font.bold = bool(negrita or bold_first)
    return marco


def _caja(slide, izq, arr, ancho, alto, lineas, **kwargs):
    caja = slide.shapes.add_textbox(izq, arr, ancho, alto)
    _texto(caja.text_frame, lineas, **kwargs)
    return caja


def _rect(slide, izq, arr, ancho, alto, color, forma=MSO_SHAPE.RECTANGLE):
    figura = slide.shapes.add_shape(forma, izq, arr, ancho, alto)
    figura.fill.solid()
    figura.fill.fore_color.rgb = color
    figura.line.fill.background()
    figura.shadow.inherit = False
    return figura


def _slide_base(prs, titulo, subtitulo=None):
    """Diapositiva de contenido: título, regla teal y pie con numeración."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    _contador["n"] += 1

    caja = slide.shapes.add_textbox(Inches(0.7), Inches(0.38), Inches(11.9), Inches(0.72))
    parrafo = caja.text_frame.paragraphs[0]
    corrida = parrafo.add_run()
    corrida.text = titulo
    corrida.font.size = Pt(30)
    corrida.font.bold = True
    corrida.font.name = FUENTE
    corrida.font.color.rgb = OSCURO

    _rect(slide, Inches(0.72), Inches(1.12), Inches(1.5), Pt(4), TEAL)

    if subtitulo:
        sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.24), Inches(11.9), Inches(0.42))
        parrafo = sub.text_frame.paragraphs[0]
        corrida = parrafo.add_run()
        corrida.text = subtitulo
        corrida.font.size = Pt(14.5)
        corrida.font.name = FUENTE
        corrida.font.color.rgb = TEAL
        corrida.font.italic = True

    pie = slide.shapes.add_textbox(Inches(11.6), Inches(6.92), Inches(1.3), Inches(0.36))
    parrafo = pie.text_frame.paragraphs[0]
    parrafo.alignment = PP_ALIGN.RIGHT
    corrida = parrafo.add_run()
    corrida.text = str(_contador["n"])
    corrida.font.size = Pt(11)
    corrida.font.name = FUENTE
    corrida.font.color.rgb = GRIS_CLARO

    pie2 = slide.shapes.add_textbox(Inches(0.7), Inches(6.92), Inches(7.0), Inches(0.36))
    parrafo = pie2.text_frame.paragraphs[0]
    corrida = parrafo.add_run()
    corrida.text = "CC3084 · Proyecto 2 · RSNA 2022 Cervical Spine Fracture Detection"
    corrida.font.size = Pt(9.5)
    corrida.font.name = FUENTE
    corrida.font.color.rgb = GRIS_CLARO
    return slide


def _figura(slide, nombre, izq, arr, ancho_max, alto_max):
    """Inserta una figura conservando su proporción y centrándola en el hueco."""
    ruta = FIGURAS / nombre
    if not ruta.exists():
        marcador = _rect(slide, izq, arr, ancho_max, alto_max, FONDO_SUAVE)
        marco = marcador.text_frame
        marco.word_wrap = True
        marco.vertical_anchor = MSO_ANCHOR.MIDDLE
        parrafo = marco.paragraphs[0]
        parrafo.alignment = PP_ALIGN.CENTER
        corrida = parrafo.add_run()
        corrida.text = f"[FALTA LA FIGURA]\n{nombre}\n\nEjecuta el notebook y expórtala a figures/"
        corrida.font.size = Pt(14)
        corrida.font.name = FUENTE
        corrida.font.color.rgb = NARANJA
        corrida.font.bold = True
        print(f"  AVISO: falta {ruta}")
        return None

    from PIL import Image  # python-pptx ya arrastra Pillow

    with Image.open(ruta) as imagen:
        ancho_px, alto_px = imagen.size
    escala = min(ancho_max / ancho_px, alto_max / alto_px)
    ancho, alto = int(ancho_px * escala), int(alto_px * escala)
    izq_centrado = izq + int((ancho_max - ancho) / 2)
    arr_centrado = arr + int((alto_max - alto) / 2)
    return slide.shapes.add_picture(str(ruta), Emu(izq_centrado), Emu(arr_centrado),
                                    width=Emu(ancho), height=Emu(alto))


def _tamano_cifra(numero: str, ancho_pulgadas: float) -> float:
    """Escoge el cuerpo de la cifra para que no se desborde de su tarjeta.

    Segoe UI Bold ocupa aproximadamente 0.58 em por carácter. Se despeja el
    cuerpo máximo que cabe en el ancho útil y se acota a un rango legible.
    """
    utiles = max(ancho_pulgadas - 0.42, 0.5)
    maximo = (utiles * 72) / (max(len(numero), 1) * 0.58)
    return max(13.0, min(34.0, maximo))


def _cifra_clave(slide, izq, arr, ancho, numero, etiqueta, color=NARANJA,
                 alto=Inches(1.72)):
    """Tarjeta con una cifra grande: es lo que se lee desde el fondo del aula."""
    _rect(slide, izq, arr, ancho, alto, FONDO_SUAVE)
    _rect(slide, izq, arr, Pt(5), alto, color)

    cuerpo = _tamano_cifra(numero, ancho / Inches(1))
    caja = slide.shapes.add_textbox(izq + Inches(0.18), arr + Inches(0.07),
                                    ancho - Inches(0.3), Inches(0.72))
    marco = caja.text_frame
    marco.word_wrap = True
    parrafo = marco.paragraphs[0]
    corrida = parrafo.add_run()
    corrida.text = numero
    corrida.font.size = Pt(cuerpo)
    corrida.font.bold = True
    corrida.font.name = FUENTE
    corrida.font.color.rgb = color

    caja2 = slide.shapes.add_textbox(izq + Inches(0.18), arr + Inches(0.8),
                                     ancho - Inches(0.3), alto - Inches(0.86))
    marco = caja2.text_frame
    marco.word_wrap = True
    parrafo = marco.paragraphs[0]
    parrafo.line_spacing = 1.1
    corrida = parrafo.add_run()
    corrida.text = etiqueta
    corrida.font.size = Pt(11)
    corrida.font.name = FUENTE
    corrida.font.color.rgb = GRIS


# ---------------------------------------------------------------------------
# Diapositivas
# ---------------------------------------------------------------------------
def construir() -> Path:
    prs = Presentation()
    prs.slide_width = ANCHO
    prs.slide_height = ALTO

    # --- 1. Portada ----------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, ANCHO, ALTO, OSCURO)
    _rect(slide, 0, Inches(6.9), ANCHO, Inches(0.6), TEAL)

    caja = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.9))
    marco = caja.text_frame
    marco.word_wrap = True
    parrafo = marco.paragraphs[0]
    corrida = parrafo.add_run()
    corrida.text = "Análisis exploratorio del conjunto\nRSNA 2022 Cervical Spine Fracture Detection"
    corrida.font.size = Pt(40)
    corrida.font.bold = True
    corrida.font.name = FUENTE
    corrida.font.color.rgb = BLANCO

    _rect(slide, Inches(1.03), Inches(4.35), Inches(2.2), Pt(5), NARANJA)
    _caja(slide, Inches(1.0), Inches(4.65), Inches(11.0), Inches(1.7), [
        ("Proyecto 2 · Reto 20", True),
        "CC3084 — Data Science · Universidad del Valle de Guatemala · Semestre II 2026",
        "Miembros A, B, C y D",
    ], size=16, color=RGBColor(0xD4, 0xDE, 0xE2), space=7)
    for forma in slide.shapes:
        if forma.has_text_frame:
            for parrafo in forma.text_frame.paragraphs:
                for corrida in parrafo.runs:
                    if corrida.font.color.rgb == OSCURO:
                        corrida.font.color.rgb = BLANCO
    _contador["n"] = 1

    # --- 2. El enunciado dice "radiografías" ---------------------------------
    slide = _slide_base(prs, "Primera corrección: no son radiografías",
                        "La unidad de análisis es un volumen 3D, no una imagen 2D")
    _caja(slide, Inches(0.72), Inches(1.85), Inches(6.4), Inches(4.6), [
        ("El enunciado del curso describe el reto como \"radiografías\".", True),
        "Es incorrecto, y la corrección cambia todo el diseño del análisis:",
        ("• Cada paciente es una tomografía computarizada (TC) axial en DICOM.", False),
        ("• Un estudio = un volumen 3D de cientos de cortes, no una imagen.", False),
        ("• La observación es el estudio, no el corte: confundirlos inflaría la", False),
        ("   muestra y filtraría información entre particiones.", False),
        ("• La evidencia diagnóstica se aprecia recorriendo el eje craneocaudal.", False),
    ], size=15, space=9)
    _cifra_clave(slide, Inches(7.5), Inches(1.95), Inches(2.5), "314", "cortes por estudio (mediana)\nrango: 69 a 1 082", TEAL)
    _cifra_clave(slide, Inches(10.25), Inches(1.95), Inches(2.5), "≈350 GB", "≈700 000 archivos DICOM\nejecución obligada en Kaggle", TEAL)
    _cifra_clave(slide, Inches(7.5), Inches(3.8), Inches(5.25), "12 instituciones · 9 países · 6 continentes",
                 "El origen multinacional explica casi toda la heterogeneidad que\nse describe más adelante.", NARANJA)

    # --- 3. Los datos --------------------------------------------------------
    slide = _slide_base(prs, "Los datos", "2 019 estudios de entrenamiento con cuatro niveles de información")
    _caja(slide, Inches(0.72), Inches(1.85), Inches(12.0), Inches(0.6), [
        ("La supervisión es piramidal: todos tienen etiqueta, muy pocos tienen localización.", True),
    ], size=16)
    tarjetas = [
        ("2 019 × 9", "train.csv\npatient_overall + C1–C7", TEAL),
        ("7 217", "bounding boxes\nsobre 235 estudios (11.64 %)", NARANJA),
        ("87", "segmentaciones .nii\nC1–C7 y T1–T12 (4.31 %)", NARANJA),
        ("2 019 × 19", "metadatos DICOM\nextraídos de las cabeceras", TEAL),
    ]
    for indice, (numero, etiqueta, color) in enumerate(tarjetas):
        _cifra_clave(slide, Inches(0.72 + indice * 3.12), Inches(2.6), Inches(2.85),
                     numero, etiqueta, color)
    _caja(slide, Inches(0.72), Inches(4.5), Inches(12.0), Inches(2.0), [
        ("El conjunto está limpio. El problema es otro.", True),
        "• De 32 chequeos de calidad, 12 verifican y salen limpios: 0 nulos, 0 duplicados,",
        "   0 UID sin imágenes, 0 huecos de numeración, 0 cajas fuera de marco.",
        "• patient_overall coincide con OR(C1..C7) en los 2 019 estudios: 0 inconsistencias.",
        "• No hay nada que imputar ni deduplicar. Los obstáculos reales son la orientación sin",
        "   normalizar, la heterogeneidad de adquisición y la supervisión espacial parcial.",
    ], size=15, space=8)

    # --- 4. Hallazgo: etiquetas ---------------------------------------------
    slide = _slide_base(prs, "Hallazgo 1 · La tarea binaria está balanceada; la de nivel no")
    _figura(slide, "01_frecuencia_etiquetas.png", Inches(0.7), Inches(1.75),
            Inches(7.2), Inches(4.9))
    _caja(slide, Inches(8.2), Inches(1.85), Inches(4.5), Inches(4.8), [
        ("47.60 %", True),
        "de los estudios son positivos (961 de 2 019). Balance cómodo: la exactitud no queda inflada por una clase mayoritaria.",
        ("El desbalance aparece un nivel más abajo:", True),
        "• C7: 393 casos (40.89 % de los positivos)",
        "• C2: 285   ·   C6: 277",
        "• C3: solo 73 casos",
        ("Razón C7 : C3 = 5.4 : 1", True),
        "Sin ponderación, un clasificador multietiqueta aprende a ignorar C3 y C4.",
    ], size=14, space=7)

    # --- 5. Hallazgo: validación contra la literatura ------------------------
    slide = _slide_base(prs, "Hallazgo 2 · La distribución reproduce la literatura clínica")
    _figura(slide, "01_correlacion_y_coocurrencia.png", Inches(0.7), Inches(1.75),
            Inches(7.4), Inches(4.9))
    _caja(slide, Inches(8.35), Inches(1.85), Inches(4.4), Inches(2.6), [
        ("66.14 % de las fracturas están en C2, C6 y C7", True),
        "(955 de 1 444 niveles fracturados; C7 es el más frecuente)",
        "• 62.4 % en el conjunto RSNA completo (Lin et al., 2023)",
        "• 63.3 % en 21 instituciones (Goldberg et al., 2001)",
        "Es una verificación externa de validez: el conjunto no está sesgado hacia un patrón atípico.",
    ], size=13.5, space=6)
    _caja(slide, Inches(8.35), Inches(4.55), Inches(4.4), Inches(2.1), [
        ("Las fracturas se agrupan entre vértebras vecinas", True),
        "Ninguna correlación phi supera 0.85, pero las más altas son siempre entre niveles adyacentes: la fuerza se disipa al segmento contiguo.",
        "→ Tratar C1–C7 como 7 problemas independientes desperdicia estructura.",
    ], size=13.5, space=6)

    # --- 6. Hallazgo: heterogeneidad ----------------------------------------
    slide = _slide_base(prs, "Hallazgo 3 · Dos protocolos distintos conviven en el conjunto")
    _figura(slide, "02_distribuciones_metadatos.png", Inches(0.7), Inches(1.75),
            Inches(5.6), Inches(4.9))
    _caja(slide, Inches(6.6), Inches(1.85), Inches(6.1), Inches(4.8), [
        ("Lo homogéneo:", True),
        "slice_thickness toma solo 10 valores (mediana 0.625 mm) y no produce ningún atípico por IQR: el único caso así del conjunto.",
        ("Lo heterogéneo:", True),
        "pixel_spacing toma 324 valores distintos (0.190–0.713 mm), siempre isotrópico en el plano. n_slices va de 69 a 1 082.",
        ("Ojo con cómo se mide z_extent:", True),
        "Como n_slices × grosor: 33.7–642.6 mm (σ = 81.3). Como |z_last − z_first| real: 48.5–334.4 mm (σ = 28.9). La aproximación sobreestima en el 99.9 % de los estudios.",
        ("La cobertura real es MÁS homogénea de lo que parecía.", True),
        "Lo que varía no es la anatomía sino el solapamiento de reconstrucción. Atípicos reales: n_slices 27 (1.34 %) · z_extent 29 (1.44 %). No se eliminan: son protocolo legítimo de 12 hospitales.",
    ], size=13, space=6)

    # --- 7. Hallazgo: orientación del eje z ----------------------------------
    slide = _slide_base(prs, "Hallazgo 4 · Nueve de cada diez volúmenes están al revés",
                        "El hallazgo con mayor potencial de daño silencioso")
    _figura(slide, "05_reporte_calidad.png", Inches(0.7), Inches(1.9),
            Inches(6.5), Inches(4.6))
    _cifra_clave(slide, Inches(7.6), Inches(1.9), Inches(5.1), "90.24 %",
                 "de los estudios (1 822 de 2 019) apila en sentido DESCENDENTE.\n"
                 "Solo 197 lo hacen en sentido ascendente.", NARANJA)
    _caja(slide, Inches(7.6), Inches(3.8), Inches(5.1), Inches(2.9), [
        ("Y la trampa es que parece correcto.", True),
        "• slice_gap = 0 en el 100 % de los estudios: la numeración de archivos NO tiene huecos.",
        "• Los 2 019 tienen ImagePositionPatient legible.",
        "Como la numeración es perfectamente contigua, ordenar por nombre de archivo se ve bien, no lanza ningún error… y procesa nueve de cada diez volúmenes invertidos.",
        ("→ Ordenar SIEMPRE por ImagePositionPatient[2], y afirmar la convención con un assert.", True),
    ], size=12.5, space=6)

    # --- 8. Hallazgo: shortcut learning -------------------------------------
    slide = _slide_base(prs, "Hallazgo 5 · Hay un riesgo de atajo, pequeño pero real")
    _figura(slide, "02_metadatos_por_clase.png", Inches(0.7), Inches(1.85),
            Inches(6.6), Inches(3.6))
    _caja(slide, Inches(7.6), Inches(1.8), Inches(5.1), Inches(4.9), [
        ("La prueba prevista no es aplicable.", True),
        "SIETE columnas valen \"Desconocido\" en el 100 % de los estudios (manufacturer, model, kernel, kvp, patient_sex, patient_age, study_desc): el conjunto fue des-identificado. El χ² da p = 1.0000, que no significa \"sin asociación\" sino \"sin datos\".",
        ("Pero las variables de adquisición sí difieren:", True),
        "n_slices:  p = 1.8 × 10⁻⁸   ·   P(pos > neg) = 0.428",
        "z_extent:  p = 3.3 × 10⁻³   ·   P(pos > neg) = 0.462",
        ("Medir bien reduce el atajo a la mitad.", True),
        "Con el z_extent aproximado el efecto parecía p = 8.5 × 10⁻¹⁰. Con el recorrido físico real cae a 3.3 × 10⁻³: la diferencia estaba en el solapamiento, no en la anatomía. n_slices pasa a ser la señal dominante.",
    ], size=13, space=6)
    _caja(slide, Inches(0.72), Inches(5.65), Inches(6.6), Inches(1.0), [
        ("→ Estratificar por n_slices discretizado, el mejor sustituto observable del sitio de origen.", True),
    ], size=13.5)

    # --- 9. Hallazgo central: tamaño de la lesión ---------------------------
    slide = _slide_base(prs, "Hallazgo 6 · La lesión ocupa el 2.4 % del corte",
                        "El resultado que más condiciona todo el diseño posterior")
    _figura(slide, "03_area_y_forma_bounding_boxes.png", Inches(0.7), Inches(2.0),
            Inches(7.4), Inches(4.5))
    _cifra_clave(slide, Inches(8.4), Inches(2.0), Inches(4.3), "2.415 %",
                 "del corte de 512 × 512 ocupa la caja (mediana).\nEn píxeles: 6 331 de 262 144.", NARANJA)
    _cifra_clave(slide, Inches(8.4), Inches(3.75), Inches(4.3), "41 : 1",
                 "relación fondo-a-lesión. Y la caja es un LÍMITE SUPERIOR:\nincluye vértebra y tejido circundante.", NARANJA)
    _caja(slide, Inches(8.4), Inches(5.5), Inches(4.3), Inches(1.2), [
        ("Consecuencia directa", True),
        "Clasificar el volumen completo obliga a la red a hallar una señal del 2 % en un campo enorme. Justifica el enfoque en dos etapas.",
    ], size=13, space=5)

    # --- 10. Hallazgo: anotación escasa --------------------------------------
    slide = _slide_base(prs, "Hallazgo 7 · La anotación espacial es escasa y no exhaustiva")
    _figura(slide, "03_solapamiento_anotaciones.png", Inches(0.7), Inches(1.9),
            Inches(6.3), Inches(4.6))
    _caja(slide, Inches(7.3), Inches(1.9), Inches(5.4), Inches(4.8), [
        ("Cobertura de la localización", True),
        "• Bounding box: 235 estudios = 11.64 % del conjunto",
        "   …y solo el 24.45 % de los casos positivos",
        "• Segmentación: 87 estudios = 4.31 %",
        "• Ambas a la vez: 40 estudios = 1.98 %",
        ("Tres de cada cuatro positivos no tienen localización.", True),
        "→ La ausencia de caja NO significa ausencia de fractura. Usar los estudios sin caja como negativos para un detector introduciría ruido de etiqueta masivo.",
        ("Los 40 estudios con ambas anotaciones son el único recurso", True),
        "para verificar que las máscaras NIfTI están bien alineadas con los volúmenes DICOM. Conviene no gastarlos como conjunto de prueba.",
    ], size=13, space=6)

    # --- 11. Evidencia visual: la ventana ------------------------------------
    slide = _slide_base(prs, "Por qué el preprocesamiento decide el resultado",
                        "Las tres imágenes son el mismo corte DICOM")
    _figura(slide, "04_comparacion_ventanas.png", Inches(0.9), Inches(1.85),
            Inches(11.5), Inches(3.5))
    _caja(slide, Inches(0.9), Inches(5.45), Inches(11.5), Inches(1.35), [
        ("El rango HU va de −1 000 (aire) a +3 000 (hueso cortical); la pantalla tiene 256 niveles.", True),
        "Sin ventana, la autoescala comprime todo el tejido intermedio y la imagen es inservible. La ventana ósea (C = 400, A = 1 800) conserva el contraste entre cortical y trabécula, que es donde se manifiesta la fractura. No es un ajuste estético: decide si la señal diagnóstica llega o no a la entrada del modelo.",
    ], size=13, space=5)

    # --- 12. Evidencia visual: positivos vs negativos ------------------------
    slide = _slide_base(prs, "Y por qué el problema es difícil",
                        "Positivos (arriba) frente a negativos (abajo), en posiciones comparables")
    _figura(slide, "04_positivos_vs_negativos.png", Inches(3.2), Inches(1.85),
            Inches(6.9), Inches(3.85))
    _caja(slide, Inches(0.9), Inches(5.75), Inches(11.5), Inches(1.1), [
        ("A simple vista, en un corte aislado, las dos filas no son distinguibles.", True),
        "La discontinuidad cortical es sutil a esta escala. La lectura clínica exige recorrer el volumen y consultar la reconstrucción sagital: un clasificador que opere corte a corte, sin contexto 3D ni localización previa, parte en desventaja.",
    ], size=13, space=5)

    # --- 13. Preprocesamiento -----------------------------------------------
    slide = _slide_base(prs, "Preprocesamiento derivado de los hallazgos",
                        "Ninguna operación elimina observaciones")
    columna_izq = [
        ("1 · Normalizar la orientación  ← la más crítica", True),
        "El 90.24 % de los volúmenes está apilado al revés. Ordenar por ImagePositionPatient[2], nunca por nombre de archivo. Reorientar las máscaras NIfTI y validar la superposición VISUALMENTE.",
        ("2 · Convertir a HU y remuestrear a 1×1×1 mm", True),
        "HU = pixel × RescaleSlope + RescaleIntercept. Sin remuestreo isotrópico, el mismo número de píxeles mide distintos milímetros según el hospital.",
        ("3 · Ventana ósea y descarte del rango de aire", True),
        "C = 400, A = 1 800 como vista primaria.",
    ]
    columna_der = [
        ("4 · Recortar a C1–C7 con las segmentaciones", True),
        "Homogeneiza el campo de visión sin descartar estudios y neutraliza el atajo por z_extent. El recorte propagado introduce un error no cuantificado en este EDA.",
        ("5 · Fijar de antemano la regla para etiquetas ambiguas", True),
        "En train no hay inconsistencias, pero el chequeo se mantiene: test público y privado pueden comportarse distinto.",
        ("6 · Partición estratificada y agrupada por paciente", True),
        "Por patient_overall, por nivel y por n_slices discretizado. La estratificación por fabricante no es posible: está des-identificada, igual que otras 6 columnas.",
    ]
    _caja(slide, Inches(0.72), Inches(1.9), Inches(5.9), Inches(4.9), columna_izq, size=12.5, space=6)
    _caja(slide, Inches(6.9), Inches(1.9), Inches(5.8), Inches(4.9), columna_der, size=12.5, space=6)

    # --- 14. Conclusiones ----------------------------------------------------
    slide = _slide_base(prs, "Conclusiones y siguientes pasos")
    _cifra_clave(slide, Inches(0.72), Inches(1.85), Inches(3.85), "2.4 %",
                 "del corte ocupa la lesión →\njustifica el enfoque en dos etapas", NARANJA)
    _cifra_clave(slide, Inches(4.75), Inches(1.85), Inches(3.85), "11.6 %",
                 "de los estudios tienen localización →\netiqueta débil, nunca verdad exhaustiva", NARANJA)
    _cifra_clave(slide, Inches(8.78), Inches(1.85), Inches(3.85), "90.2 %",
                 "de los volúmenes está apilado al revés →\nnormalizar orientación, no limpiar", NARANJA)
    _caja(slide, Inches(0.72), Inches(3.7), Inches(12.0), Inches(3.05), [
        ("1. Adoptar un enfoque en dos etapas: localizar la vértebra, luego clasificar el recorte.", True),
        "   Reduce la relación fondo-a-lesión, homogeneiza el campo de visión y convierte un problema multietiqueta en siete decisiones locales comparables.",
        ("2. Tratar cajas y segmentaciones como etiqueta débil.", True),
        "   Los estudios sin anotación no son negativos; un detector entrenado así aprendería a no detectar.",
        ("3. Reservar los 40 estudios con ambas anotaciones para validar el pipeline geométrico.", True),
        ("4. Fijar el preprocesamiento antes de experimentar.", True),
        "   Los errores de orientación no lanzan excepción: degradan el desempeño de forma inexplicable.",
        ("5. Medir la geometría con la cabecera, no con aritmética sobre el conteo de cortes.", True),
        "   |z_last − z_first| frente a n_slices × grosor cambió el rango, los atípicos (84→29) y la fuerza del atajo.",
    ], size=12.5, space=5)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SALIDA)
    print(f"Presentación generada: {SALIDA}")
    print(f"Diapositivas: {len(prs.slides.__iter__.__self__._sldIdLst)}  "
          f"({SALIDA.stat().st_size / 1024:.0f} KB)")
    return SALIDA


if __name__ == "__main__":
    construir()
